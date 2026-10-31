"""Smoke tests for the main window. They are skipped on a machine with no screen."""
import threading
import time

import pytest

pytest.importorskip("tkinter")

from siphon_app.core.fetcher import FetchCancelled, FetchedPage, FetchError
from siphon_app.core.models import Table
from siphon_app.gui import app as app_module
from siphon_app.gui.app import LIGHT_THEME, READY_TEXT, SiphonApp

# helpers: fake downloads, so no test ever touches the internet
PAGE = FetchedPage(url="https://example.com/data", content=b"<html></html>", content_type="text/html")

def wait_until_idle(app, timeout=5.0):
    """Let the window process events until the background job has been handled."""
    end = time.time() + timeout
    while app.busy and time.time() < end:
        app.update()
        time.sleep(0.01)
    app.update()
    assert not app.busy, "the window never finished the job"

def run_fetch(app, monkeypatch, url="example.com/data", *, fetch=None, parse=None):
    """Put a url in the box, swap in fake fetch/parse functions, and press Fetch."""
    monkeypatch.setattr(app_module, "fetch", fetch or (lambda u, cancel_event=None: PAGE))
    monkeypatch.setattr(app_module, "parse_page", parse or (lambda page: []))
    app.url_var.set(url)
    app.on_fetch()

# the window itself
def test_window_starts_ready_with_the_light_theme(app):
    assert app.title() == "siphon."
    assert app.status_var.get() == READY_TEXT
    assert app.style.theme.name == LIGHT_THEME
    assert app.style.theme_mode == "light"
    assert not app.status_is_error
    assert not app.busy and str(app.cancel_button.cget("state")) == "disabled"

def test_a_bad_address_shows_the_reason_in_the_status_bar(app):
    app.url_var.set("mailto:someone@example.com")
    app.on_fetch()
    assert "http and https" in app.status_var.get()
    assert app.status_is_error
    assert not app.busy # a bad address never starts a download

def test_an_empty_address_is_rejected(app):
    app.url_var.set("   ")
    app.on_fetch()
    assert app.status_var.get() != READY_TEXT
    assert app.status_is_error
    assert not app.busy

def test_theme_button_switches_to_dark_and_back(app):
    app.toggle_theme()
    assert app.style.theme_mode == "dark"
    assert app.theme_button.cget("text") == "Light" # the button now offers the other theme
    app.toggle_theme()
    assert app.style.theme_mode == "light"
    assert app.theme_button.cget("text") == "Dark"

def test_enter_in_the_url_box_fetches(app):
    assert app.url_entry.bind("<Return>") # a handler is attached

# downloading in the background
def test_a_good_address_is_tidied_and_downloaded(app, monkeypatch):
    # replaces the old Step 6 test: fetching now really runs (with a fake download)
    seen = []
    def fake_fetch(url, cancel_event=None):
        seen.append(url)
        return PAGE
    tables = [Table("Prices", ["A", "B"], [["1", "2"], ["3", "4"]])]
    run_fetch(app, monkeypatch, "  example.com/data ", fetch=fake_fetch, parse=lambda page: tables)
    wait_until_idle(app)
    assert app.url_var.get() == "https://example.com/data" # https:// added, spaces gone
    assert seen == ["https://example.com/data"]
    assert app.tables == tables
    assert app.status_var.get() == "Found 1 table."
    assert not app.status_is_error
    assert app.table_view.name_label.cget("text") == "Prices" # shown in the preview now
    assert "2 rows" in app.table_view.size_label.cget("text")

def test_several_tables_and_layout_tables_are_described(app, monkeypatch):
    tables = [Table("Sales", ["A", "B"], [["1", "2"]]),
            Table("Menu", ["A"], [["x"]], is_layout=True)]
    run_fetch(app, monkeypatch, parse=lambda page: tables)
    wait_until_idle(app)
    assert app.status_var.get() == "Found 2 tables (1 looks like page layout)."
    listed = [app.table_view.table_list.set(i, "name") for i in app.table_view.table_list.get_children()]
    assert listed == ["Sales"] # the layout table is hidden by default

def test_the_window_stays_responsive_and_locks_the_buttons_while_downloading(app, monkeypatch):
    started, gate = threading.Event(), threading.Event()
    def slow_fetch(url, cancel_event=None):
        started.set()
        gate.wait(5) # pretend the server is slow
        return PAGE
    run_fetch(app, monkeypatch, fetch=slow_fetch)
    assert started.wait(2)
    app.update() # the window can still process events
    assert app.busy
    assert str(app.fetch_button.cget("state")) == "disabled"
    assert str(app.cancel_button.cget("state")) == "normal"
    assert app.progress.winfo_manager() == "grid" # the progress bar is showing
    assert "Downloading from example.com" in app.status_var.get()
    gate.set()
    wait_until_idle(app)
    assert str(app.fetch_button.cget("state")) == "normal"
    assert str(app.cancel_button.cget("state")) == "disabled"
    assert app.progress.winfo_manager() == "" # and is hidden again

def test_pressing_fetch_again_while_busy_does_not_start_a_second_download(app, monkeypatch):
    calls, gate = [], threading.Event()
    def slow_fetch(url, cancel_event=None):
        calls.append(url)
        gate.wait(5)
        return PAGE
    run_fetch(app, monkeypatch, fetch=slow_fetch)
    app.on_fetch()
    app.on_fetch()
    gate.set()
    wait_until_idle(app)
    assert len(calls) == 1

def test_cancel_frees_the_window_at_once_and_a_late_answer_is_ignored(app, monkeypatch):
    stop_seen = threading.Event()
    def stubborn_fetch(url, cancel_event=None):
        cancel_event.wait(5) # waits until it is told to stop...
        stop_seen.set()
        return PAGE  # ...but even so, hands back a page
    tables_before = list(app.tables)
    run_fetch(app, monkeypatch, fetch=stubborn_fetch, parse=lambda page: [Table("Late", ["A"], [["1"]])])
    app.update()
    app.on_cancel()
    assert not app.busy # free immediately, no waiting for the worker
    assert app.status_var.get() == "Cancelled."
    assert stop_seen.wait(2) # the worker really was told to stop
    for _ in range(20): # give any late message time to arrive
        app.update()
        time.sleep(0.01)
    assert app.status_var.get() == "Cancelled." # the stale answer did not overwrite it
    assert app.tables == tables_before

def test_cancel_when_nothing_is_running_does_nothing(app):
    before = app.status_var.get()
    app.on_cancel()
    assert app.status_var.get() == before

def test_a_download_can_be_cancelled_by_the_fetcher_itself(app, monkeypatch):
    def cancelled_fetch(url, cancel_event=None):
        raise FetchCancelled("Cancelled.")
    run_fetch(app, monkeypatch, fetch=cancelled_fetch)
    wait_until_idle(app)
    assert app.status_var.get() == "Cancelled."
    assert not app.status_is_error

def test_a_failed_download_shows_the_message_in_red(app, monkeypatch):
    def failing_fetch(url, cancel_event=None):
        raise FetchError("Page not found (404). Check the address.")
    run_fetch(app, monkeypatch, fetch=failing_fetch)
    wait_until_idle(app)
    assert app.status_var.get() == "Page not found (404). Check the address."
    assert app.status_is_error

def test_a_crash_in_the_parser_is_reported_not_swallowed_and_never_hangs(app, monkeypatch):
    def broken_parse(page):
        raise RuntimeError("boom")
    run_fetch(app, monkeypatch, parse=broken_parse)
    wait_until_idle(app)
    assert "Something went wrong" in app.status_var.get() and "RuntimeError" in app.status_var.get()
    assert app.status_is_error

def test_a_page_without_tables_says_so_and_clears_the_old_list(app, monkeypatch):
    run_fetch(app, monkeypatch, parse=lambda page: [Table("Old", ["A"], [["1"]])])
    wait_until_idle(app)
    run_fetch(app, monkeypatch, parse=lambda page: [])
    wait_until_idle(app)
    assert app.tables == []
    assert "No tables found" in app.status_var.get()
    assert not app.status_is_error
    assert app.table_view.checked_tables() == []                     # [CHANGED]
    assert app.table_view.placeholder.cget("text") == "No tables found on that page."

def test_a_new_download_works_after_an_error(app, monkeypatch):
    run_fetch(app, monkeypatch, fetch=lambda u, cancel_event=None: (_ for _ in ()).throw(FetchError("nope")))
    wait_until_idle(app)
    run_fetch(app, monkeypatch, parse=lambda page: [Table("Ok", ["A"], [["1"]])])
    wait_until_idle(app)
    assert app.status_var.get() == "Found 1 table."
    assert not app.status_is_error

def test_an_answer_from_a_cancelled_job_cannot_leak_into_the_next_one(app, monkeypatch):
    # The cancelled worker still sends its answer later. It must not be shown, and it
    # must not end the NEXT download early, when that download's poll picks it up.
    release_old, release_new = threading.Event(), threading.Event()
    def late_fetch(url, cancel_event=None):
        release_old.wait(5)
        return FetchedPage(url=url, content=b"Stale", content_type="text/html")
    def fresh_fetch(url, cancel_event=None):
        release_new.wait(5)
        return FetchedPage(url=url, content=b"Fresh", content_type="text/html")
    def parse(page): # the table is named after the page it came from
        return [Table(page.content.decode(), ["A"], [["1"]])]
    run_fetch(app, monkeypatch, fetch=late_fetch, parse=parse)
    app.update()
    app.on_cancel()
    release_old.set()
    time.sleep(0.3) # the old worker has now queued its answer
    run_fetch(app, monkeypatch, fetch=fresh_fetch, parse=parse)
    for _ in range(20): # the poll picks up the stale answer here
        app.update()
        time.sleep(0.01)
    assert app.busy # the stale answer did not end the new job
    assert [t.name for t in app.tables] != ["Stale"]
    release_new.set()
    wait_until_idle(app)
    assert [t.name for t in app.tables] == ["Fresh"]
    assert app.table_view.name_label.cget("text") == "Fresh"

def test_run_job_reports_cancelled_if_the_stop_signal_came_while_the_page_arrived(monkeypatch):
    # [NEW] The worker function can be tested without any window at all.
    import queue
    monkeypatch.setattr(app_module, "fetch", lambda url, cancel_event=None: PAGE)
    monkeypatch.setattr(app_module, "parse_page", lambda page: pytest.fail("should not parse after a cancel"))
    results, event = queue.Queue(), threading.Event()
    event.set()
    app_module.run_job(results, 7, "https://example.com", event)
    assert results.get_nowait() == ("cancelled", 7)

def test_cancel_tells_the_worker_to_stop(app, monkeypatch):
    #the worker is only told through the cancel event, so check it really arrives
    seen = {}
    got_event = threading.Event()
    def fetch_that_records(url, cancel_event=None):
        seen["event"] = cancel_event
        got_event.set()
        cancel_event.wait(5)
        raise FetchCancelled("Cancelled.")
    run_fetch(app, monkeypatch, fetch=fetch_that_records)
    assert got_event.wait(2)
    assert not seen["event"].is_set()
    app.on_cancel()
    assert seen["event"].is_set()

#cleaning and exporting through window
import csv

from siphon_app.gui.export_bar import CLEAN_OFF_TEXT


def prices():
    return [Table("Prices", ["Name", "Price"], [["Alpha", "$1,234"], ["Beta", "$5"]],
                source_url="https://example.com/data")]

def preview_rows(app):
    view = app.table_view.preview
    return [[str(v) for v in view.item(i, "values")] for i in view.get_children()]

def wait_for_export(app, timeout=5.0):
    end = time.time() + timeout
    while app.export_bar.busy and time.time() < end:
        app.update()
        time.sleep(0.01)
    app.update()
    assert not app.export_bar.busy, "the export never finished"

def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))

@pytest.fixture(autouse=True)
def reset_export_bar(app):
    """The window is shared by every test, so put the export controls back afterwards."""
    yield
    app.export_bar.clean_var.set(True)
    app.export_bar.format_var.set("csv")

def test_numbers_are_cleaned_by_default_and_the_box_flips_between_before_and_after(app, monkeypatch):
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    assert app.export_bar.clean_var.get() is True
    assert preview_rows(app) == [["Alpha", "1234"], ["Beta", "5"]] # after
    assert "Turned 2 cells in 1 column into numbers." in app.export_bar.clean_label.cget("text")

    app.export_bar.clean_var.set(False)
    app.export_bar._on_clean()
    assert preview_rows(app) == [["Alpha", "$1,234"], ["Beta", "$5"]] # before
    assert app.export_bar.clean_label.cget("text") == CLEAN_OFF_TEXT

    app.export_bar.clean_var.set(True)
    app.export_bar._on_clean()
    assert preview_rows(app) == [["Alpha", "1234"], ["Beta", "5"]]
    assert "Turned 2 cells" in app.export_bar.clean_label.cget("text")

def test_a_new_page_respects_the_clean_box_being_off(app, monkeypatch):
    app.export_bar.clean_var.set(False)
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    assert preview_rows(app) == [["Alpha", "$1,234"], ["Beta", "$5"]]
    assert app.export_bar.clean_label.cget("text") == CLEAN_OFF_TEXT

def test_toggling_the_clean_box_keeps_the_ticks_and_the_chosen_table(app, monkeypatch):
    tables = [Table("One", ["A"], [["$1"]]), Table("Two", ["A"], [["$2"]])]
    run_fetch(app, monkeypatch, parse=lambda page: tables)
    wait_until_idle(app)
    app.table_view.toggle(0) # untick One
    app.table_view.table_list.selection_set("1")
    app.update()
    app.export_bar.clean_var.set(False)
    app.export_bar._on_clean()
    assert app.table_view.name_label.cget("text") == "Two"
    assert [t.name for t in app.table_view.checked_tables()] == ["Two"]
    assert app.table_view.checked_tables()[0].rows == [["$2"]] # the original text

def test_toggling_with_no_tables_is_harmless(app, monkeypatch):
    run_fetch(app, monkeypatch, parse=lambda page: [])
    wait_until_idle(app)
    app.export_bar.clean_var.set(False)
    app.export_bar._on_clean()
    assert app.export_bar.clean_label.cget("text") == CLEAN_OFF_TEXT
    app.export_bar.clean_var.set(True)
    app.export_bar._on_clean()
    assert "Numbers like" in app.export_bar.clean_label.cget("text")

def test_cleaning_happens_on_the_worker_thread_not_in_the_window(app, monkeypatch):
    seen = []
    real = app_module.clean_all
    monkeypatch.setattr(app_module, "clean_all",
                        lambda tables: seen.append(threading.current_thread()) or real(tables))
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    assert seen and seen[0] is not threading.main_thread()

def test_the_export_button_follows_the_ticks(app, monkeypatch):
    tables = [Table("Sales", ["A"], [["1"]]), Table("Menu", ["A"], [["x"]], is_layout=True)]
    run_fetch(app, monkeypatch, parse=lambda page: tables)
    wait_until_idle(app)
    assert app.export_bar.export_button.cget("text") == "Export 1 table"  # layout is unticked
    app.table_view.select_none()
    assert app.export_bar.export_button.cget("text") == "Export"
    assert str(app.export_bar.export_button.cget("state")) == "disabled"
    app.table_view.select_all()
    assert str(app.export_bar.export_button.cget("state")) == "normal"
    app.table_view.hide_layout.set(False)
    app.table_view._refresh_list()
    app.table_view.select_all()
    assert app.export_bar.export_button.cget("text") == "Export 2 tables"

def test_the_export_button_is_off_when_a_page_has_no_tables(app, monkeypatch):
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    run_fetch(app, monkeypatch, parse=lambda page: [])
    wait_until_idle(app)
    assert str(app.export_bar.export_button.cget("state")) == "disabled"
    assert app.export_bar.export_button.cget("text") == "Export"

def test_exporting_from_the_window_saves_what_the_preview_shows(app, monkeypatch, tmp_path):
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    app.export_bar.folder_var.set(str(tmp_path))
    app.export_bar.on_export()
    wait_for_export(app)
    assert app.status_var.get() == f"Saved Prices.csv to {tmp_path}."
    assert not app.status_is_error
    assert read_csv(tmp_path / "Prices.csv") == [["Name", "Price"], ["Alpha", "1234"], ["Beta", "5"]]

    app.export_bar.clean_var.set(False) # now the untouched text
    app.export_bar._on_clean()
    app.export_bar.on_export()
    wait_for_export(app)
    assert read_csv(tmp_path / "Prices (2).csv")[1] == ["Alpha", "$1,234"]
    assert read_csv(tmp_path / "Prices.csv")[1] == ["Alpha", "1234"] # the first file is safe

def test_only_ticked_tables_are_exported(app, monkeypatch, tmp_path):
    tables = [Table("One", ["A"], [["1"]]), Table("Two", ["A"], [["2"]]), Table("Three", ["A"], [["3"]])]
    run_fetch(app, monkeypatch, parse=lambda page: tables)
    wait_until_idle(app)
    app.table_view.toggle(1) # untick Two
    app.export_bar.folder_var.set(str(tmp_path))
    app.export_bar.on_export()
    wait_for_export(app)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["One.csv", "Three.csv"]
    assert app.status_var.get() == f"Saved 2 files to {tmp_path}."

def test_an_excel_export_from_the_window_has_real_numbers(app, monkeypatch, tmp_path):
    from openpyxl import load_workbook
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    app.export_bar.format_var.set("xlsx")
    app.export_bar.folder_var.set(str(tmp_path))
    app.export_bar.on_export()
    wait_for_export(app)
    assert load_workbook(tmp_path / "Prices.xlsx")["Prices"]["B2"].value == 1234

def test_a_failed_export_shows_in_red_in_the_status_bar(app, monkeypatch, tmp_path):
    blocker = tmp_path / "file.txt"
    blocker.write_text("hi")
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    app.export_bar.folder_var.set(str(blocker))
    app.export_bar.on_export()
    wait_for_export(app)
    assert app.status_is_error and "Couldn't" in app.status_var.get()

#remembering things and opening saved
from siphon_app.core.parser import parse_page as real_parse_page
from siphon_app.gui.export_bar import default_folder
from siphon_app.settings import Settings, load_settings

SAVED_PAGE = (b"<html><body><h2>Sales</h2><table><tr><th>Name</th><th>Total</th></tr>"
            b"<tr><td>A</td><td>$5</td></tr><tr><td>B</td><td>$7</td></tr><tr><td>C</td><td>$9</td></tr>"
            b"</table></body></html>")

@pytest.fixture
def own_settings(app, tmp_path):
    """Give the shared window a fresh, private settings file for one test, then put things back."""
    old = (app.settings, app.settings_file, list(app.url_entry.cget("values")))
    app.settings = Settings()
    app.settings_file = tmp_path / "settings.json"
    app.url_entry.configure(values=[])
    yield app.settings_file
    app.settings, app.settings_file = old[0], old[1]
    app.url_entry.configure(values=old[2])
    if app.style.theme_mode == "dark":
        app.toggle_theme()

def saved_html(tmp_path, name="page.html", content=SAVED_PAGE):
    path = tmp_path / name
    path.write_bytes(content)
    return path

def choose_file(monkeypatch, result):
    """Pretend the person picked `result` in the Open dialog ('' means they pressed Cancel)."""
    monkeypatch.setattr(app_module.filedialog, "askopenfilename", lambda **kwargs: str(result))

#recent pages
def test_a_successful_fetch_is_remembered_and_saved(app, monkeypatch, own_settings):
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    assert app.settings.recent_urls == ["https://example.com/data"]
    assert list(app.url_entry.cget("values")) == ["https://example.com/data"]
    assert load_settings(own_settings).recent_urls == ["https://example.com/data"] # really on disk

def test_a_failed_fetch_is_not_remembered(app, monkeypatch, own_settings):
    def broken(url, cancel_event=None):
        raise FetchError("nope")
    run_fetch(app, monkeypatch, fetch=broken)
    wait_until_idle(app)
    assert app.settings.recent_urls == [] and not own_settings.exists()

def test_a_cancelled_fetch_is_not_remembered(app, monkeypatch, own_settings):
    release = threading.Event()
    def slow(url, cancel_event=None):
        release.wait(2)
        return PAGE
    run_fetch(app, monkeypatch, fetch=slow)
    app.on_cancel()
    release.set()
    time.sleep(0.1)
    app.update()
    assert app.settings.recent_urls == []

def test_fetching_the_same_page_twice_lists_it_once_and_newest_goes_first(app, monkeypatch, own_settings):
    for url in ("a.com/x", "b.com/y", "a.com/x"):
        run_fetch(app, monkeypatch, url=url)
        wait_until_idle(app)
    assert app.settings.recent_urls == ["https://a.com/x", "https://b.com/y"]

def test_picking_a_recent_page_from_the_list_fetches_it(app, monkeypatch, own_settings):
    seen = []
    monkeypatch.setattr(app_module, "fetch", lambda u, cancel_event=None: seen.append(u) or PAGE)
    monkeypatch.setattr(app_module, "parse_page", lambda page: [])
    app.url_var.set("https://example.com/older")
    app.url_entry.event_generate("<<ComboboxSelected>>")
    wait_until_idle(app)
    assert seen == ["https://example.com/older"]

#the saved choices
def test_collect_settings_reads_the_controls(app, own_settings, tmp_path):
    app.export_bar.folder_var.set(f"  {tmp_path}  ")
    app.export_bar.format_var.set("json")
    app.export_bar.clean_var.set(False)
    app.toggle_theme()
    got = app.collect_settings()
    assert (got.export_folder, got.export_format, got.clean_numbers, got.theme) == (str(tmp_path), "json", False, "dark")

def test_collect_settings_is_a_copy_and_changes_nothing(app, own_settings):
    app.settings.recent_urls.append("https://a.com")
    app.export_bar.format_var.set("xlsx")
    got = app.collect_settings()
    got.recent_urls.append("https://b.com")
    assert app.settings.recent_urls == ["https://a.com"] and app.settings.export_format == "csv"

def test_a_hidden_window_keeps_the_remembered_size(app, own_settings):
    app.settings.window_width, app.settings.window_height = 1234, 777
    got = app.collect_settings()                       # the test window is withdrawn, so its size is meaningless
    assert (got.window_width, got.window_height) == (1234, 777)

def test_save_settings_writes_the_file(app, own_settings):
    app.export_bar.format_var.set("xlsx")
    app.save_settings()
    assert load_settings(own_settings).export_format == "xlsx"
    assert app.settings.export_format == "xlsx"

def test_a_save_that_fails_does_not_bother_the_window(app, own_settings, tmp_path):
    blocker = tmp_path / "file.txt"
    blocker.write_text("x")
    app.settings_file = blocker / "settings.json"
    app.save_settings()  # must not raise

def test_a_successful_export_remembers_the_folder_and_format(app, monkeypatch, own_settings, tmp_path):
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    out = tmp_path / "out"
    out.mkdir()
    app.export_bar.folder_var.set(str(out))
    app.export_bar.format_var.set("xlsx")
    app.export_bar.clean_var.set(False)
    app.export_bar._on_clean()
    app.export_bar.on_export()
    wait_for_export(app)
    saved = load_settings(own_settings)
    assert (saved.export_folder, saved.export_format, saved.clean_numbers) == (str(out), "xlsx", False)

def test_a_failed_export_does_not_save_anything(app, monkeypatch, own_settings, tmp_path):
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    own_settings.unlink(missing_ok=True)
    blocker = tmp_path / "file.txt"
    blocker.write_text("x")
    app.export_bar.folder_var.set(str(blocker))
    app.export_bar.on_export()
    wait_for_export(app)
    assert not own_settings.exists() or load_settings(own_settings).export_folder != str(blocker)

def test_remembered_choices_are_put_back_into_the_controls(app, own_settings, tmp_path):
    app.settings = Settings(export_folder=str(tmp_path), export_format="json", clean_numbers=False, theme="dark")
    app._apply_settings()
    bar = app.export_bar
    assert (bar.folder_var.get(), bar.format_var.get(), bar.clean_var.get()) == (str(tmp_path), "json", False)
    assert bar.clean_label.cget("text") == CLEAN_OFF_TEXT
    assert app.theme_button.cget("text") == "Light"
    app.export_bar.format_var.set("csv")
    app.export_bar.clean_var.set(True)
    app._apply_settings()                              # restore the label text for later tests
    app.settings = Settings()
    app._apply_settings()

def test_a_remembered_folder_that_vanished_falls_back_to_downloads(app, own_settings, tmp_path):
    app.settings = Settings(export_folder=str(tmp_path / "deleted"))
    app._apply_settings()
    assert app.export_bar.folder_var.get() == default_folder()
    app.settings = Settings()
    app._apply_settings()

#opening a saved page
def test_open_file_shows_the_tables_and_leaves_the_address_box_and_recents_alone(app, monkeypatch, own_settings, tmp_path):
    monkeypatch.setattr(app_module, "parse_page", real_parse_page)
    choose_file(monkeypatch, saved_html(tmp_path))
    app.url_var.set("https://typed.example/")
    app.on_open_file()
    assert app.busy and "Reading page.html" in app.status_var.get()
    wait_until_idle(app)
    assert [t.name for t in app.tables] == ["Sales"]
    assert app.status_var.get().startswith("Found 1 table")
    assert app.url_var.get() == "https://typed.example/"
    assert app.settings.recent_urls == [] and not own_settings.exists()

def test_open_file_then_export_works_like_any_page(app, monkeypatch, own_settings, tmp_path):
    monkeypatch.setattr(app_module, "parse_page", real_parse_page)
    choose_file(monkeypatch, saved_html(tmp_path))
    app.on_open_file()
    wait_until_idle(app)
    out = tmp_path / "out"
    out.mkdir()
    app.export_bar.folder_var.set(str(out))
    app.export_bar.on_export()
    wait_for_export(app)
    assert read_csv(out / "Sales.csv") == [["Name", "Total"], ["A", "5"], ["B", "7"], ["C", "9"]]

def test_cancelling_the_open_dialog_does_nothing(app, monkeypatch, own_settings):
    app.set_status("before")
    choose_file(monkeypatch, "")
    app.on_open_file()
    assert not app.busy and app.status_var.get() == "before"

def test_a_file_that_is_not_a_web_page_shows_a_red_message(app, monkeypatch, own_settings, tmp_path):
    choose_file(monkeypatch, saved_html(tmp_path, "notes.txt"))
    app.on_open_file()
    wait_until_idle(app)
    assert app.status_is_error and "saved web pages" in app.status_var.get()

def test_a_missing_file_shows_a_red_message(app, monkeypatch, own_settings, tmp_path):
    choose_file(monkeypatch, tmp_path / "gone.html")
    app.on_open_file()
    wait_until_idle(app)
    assert app.status_is_error and "doesn't exist" in app.status_var.get()

def test_open_file_is_ignored_while_a_job_is_running(app, monkeypatch, own_settings):
    asked = []
    monkeypatch.setattr(app_module.filedialog, "askopenfilename", lambda **kw: asked.append(1) or "")
    app.busy = True
    try:
        app.on_open_file()
    finally:
        app.busy = False
    assert asked == []

def test_the_open_button_is_disabled_while_busy_and_ctrl_o_is_bound(app, monkeypatch):
    release = threading.Event()
    run_fetch(app, monkeypatch, fetch=lambda u, cancel_event=None: release.wait(2) or PAGE)
    assert str(app.open_button.cget("state")) == "disabled"
    app.on_cancel()
    release.set()
    assert str(app.open_button.cget("state")) == "normal"
    assert app.bind("<Control-o>")

def test_the_tests_never_touch_the_real_settings_file(app):
    from siphon_app.settings import settings_path
    assert app.settings_file != settings_path()

#guards that will show up in unusual situations
def test_opening_a_file_raises_no_hidden_errors(app, monkeypatch, own_settings, tmp_path):
    """Errors inside window callbacks are swallowed by Tk, so catch them explicitly."""
    problems = []
    monkeypatch.setattr(app, "report_callback_exception", lambda *info: problems.append(info))
    monkeypatch.setattr(app_module, "parse_page", real_parse_page)
    choose_file(monkeypatch, saved_html(tmp_path))
    app.on_open_file()
    wait_until_idle(app)
    assert problems == []

def test_a_maximised_window_does_not_overwrite_the_remembered_size(app, own_settings, monkeypatch):
    app.settings.window_width, app.settings.window_height = 1000, 640
    monkeypatch.setattr(app, "winfo_width", lambda: 1900)
    monkeypatch.setattr(app, "winfo_height", lambda: 1000)
    monkeypatch.setattr(app, "state", lambda *a: "zoomed")
    got = app.collect_settings()
    assert (got.window_width, got.window_height) == (1000, 640)
    monkeypatch.setattr(app, "state", lambda *a: "normal")
    got = app.collect_settings()
    assert (got.window_width, got.window_height) == (1900, 1000)

def test_closing_the_window_saves_the_settings(app, own_settings, monkeypatch):
    monkeypatch.setattr(app, "destroy", lambda: None) # keep the shared test window alive
    app.export_bar.format_var.set("json")
    app.on_close()
    assert load_settings(own_settings).export_format == "json"

def test_a_note_about_a_remembered_copy_is_added_to_the_status(app, monkeypatch, own_settings):
    noted = FetchedPage(url="https://example.com/data", content=b"<html></html>", content_type="text/html",
                        note="Used the copy from 40 seconds ago; the site wasn't contacted.")
    run_fetch(app, monkeypatch, fetch=lambda u, cancel_event=None: noted, parse=lambda page: prices())
    wait_until_idle(app)
    text = app.status_var.get()
    assert text.startswith("Found 1 table.") and text.endswith("the site wasn't contacted.")
    assert not app.status_is_error

def test_no_note_means_the_status_is_unchanged(app, monkeypatch, own_settings):
    run_fetch(app, monkeypatch, parse=lambda page: prices())
    wait_until_idle(app)
    assert app.status_var.get() == "Found 1 table."

def test_a_site_that_says_no_robots_shows_its_reason_in_red(app, monkeypatch, own_settings):
    def refused(url, cancel_event=None):
        raise FetchError("example.com asks robots not to read this page (robots.txt), so siphon didn't fetch it.")
    run_fetch(app, monkeypatch, fetch=refused)
    wait_until_idle(app)
    assert app.status_is_error and "robots.txt" in app.status_var.get()
    assert app.settings.recent_urls == []

def test_reload_forgets_the_remembered_copy_first(app, monkeypatch, own_settings):
    forgotten = []
    monkeypatch.setattr(app_module, "forget", forgotten.append)
    monkeypatch.setattr(app_module, "fetch", lambda u, cancel_event=None: PAGE)
    monkeypatch.setattr(app_module, "parse_page", lambda page: [])
    app.url_var.set("example.com/data")
    app.on_fetch(fresh=True)
    wait_until_idle(app)
    assert forgotten == ["https://example.com/data"]

def test_a_normal_fetch_keeps_the_remembered_copy(app, monkeypatch, own_settings):
    forgotten = []
    monkeypatch.setattr(app_module, "forget", forgotten.append)
    run_fetch(app, monkeypatch)
    wait_until_idle(app)
    assert forgotten == []

def test_reload_with_a_bad_address_forgets_nothing(app, monkeypatch):
    forgotten = []
    monkeypatch.setattr(app_module, "forget", forgotten.append)
    app.url_var.set("mailto:someone@example.com")
    app.on_fetch(fresh=True)
    assert forgotten == [] and app.status_is_error and not app.busy

def test_ctrl_r_is_bound_to_reload(app):
    assert app.bind("<Control-r>")

def test_the_window_downloads_through_the_polite_fetcher():
    from siphon_app.core import polite
    assert app_module.fetch is polite.polite_fetch