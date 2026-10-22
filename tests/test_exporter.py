"""Tests for siphon_app/core/exporter.py. Run them with:  pytest"""
import csv
import json

import pytest
from openpyxl import load_workbook

from siphon_app.core import exporter
from siphon_app.core.exporter import (
    ExportError, export_csv, export_json, export_xlsx, safe_filename, safe_sheet_names,
)
from siphon_app.core.models import Table

def sample(name="Prices"):
    return Table(name, ["Name", "Price"], [["Alpha", 1234.5], ["Beta", 99], ["Gamma", ""]],
                source_url="http://x.test/p")

#names
@pytest.mark.parametrize("raw, expected", [
    ("Population by country", "Population by country"),
    ('A/B\\C:D*E?F"G<H>I|J', "A B C D E F G H I J"),
    ("  spaced   out  ", "spaced out"),
    ("ends with dots...", "ends with dots"),
    ("", "table"), ("???", "table"), ("...", "table"),
    ("con", "_con"), ("NUL", "_NUL"), ("com1", "_com1"), ("lpt9.backup", "_lpt9.backup"),
    ("console", "console"),                                   # only the exact reserved words
    ("tab\there\nnewline", "tab here newline"),
])
def test_safe_filename(raw, expected):
    assert safe_filename(raw) == expected

def test_safe_filename_is_shortened():
    assert len(safe_filename("x" * 500)) == 100

def test_safe_sheet_names():
    names = safe_sheet_names(["Sales [2020]: Q1/Q2?", "x" * 50, "", "'quoted'", "History", "history"])
    assert names[0] == "Sales 2020 Q1 Q2"
    assert names[1] == "x" * 31
    assert names[2] == "Sheet"
    assert names[3] == "quoted"
    assert names[4] == "History_" and names[5] == "history_ (2)"
    assert all(len(n) <= 31 for n in names)

def test_sheet_names_are_unique_ignoring_capitals_and_stay_short():
    names = safe_sheet_names(["Stats", "STATS", "stats", "y" * 40, "y" * 40])
    assert names == ["Stats", "STATS (2)", "stats (3)", "y" * 31, "y" * 27 + " (2)"]

#CSV
def test_csv_has_a_bom_and_the_right_content(tmp_path):
    (path,) = export_csv([sample()], tmp_path)
    assert path.name == "Prices.csv"
    assert path.read_bytes().startswith(b"\xef\xbb\xbf") # so Excel reads accents right
    with open(path, encoding="utf-8-sig", newline="") as handle:
        assert list(csv.reader(handle)) == [["Name", "Price"], ["Alpha", "1234.5"], ["Beta", "99"], ["Gamma", ""]]

def test_csv_keeps_non_english_text(tmp_path):
    table = Table("t", ["City"], [["\u0141\u00f3d\u017a"], ["\u6771\u4eac"], ["caf\u00e9, \"quoted\"\nline"]])
    (path,) = export_csv([table], tmp_path)
    with open(path, encoding="utf-8-sig", newline="") as handle:
        assert [r[0] for r in list(csv.reader(handle))[1:]] == ["\u0141\u00f3d\u017a", "\u6771\u4eac", "caf\u00e9, \"quoted\"\nline"]

def test_csv_one_file_per_table_with_unique_safe_names(tmp_path):
    tables = [sample("Stats"), sample("stats"), sample("a/b:c"), sample("CON")]
    paths = export_csv(tables, tmp_path)
    assert [p.name for p in paths] == ["Stats.csv", "stats (2).csv", "a b c.csv", "_CON.csv"]
    assert all(p.exists() for p in paths)

def test_csv_never_overwrites_an_existing_file(tmp_path):
    (tmp_path / "Prices.csv").write_text("precious")
    (path,) = export_csv([sample()], tmp_path)
    assert path.name == "Prices (2).csv"
    assert (tmp_path / "Prices.csv").read_text() == "precious"

def test_csv_creates_the_folder(tmp_path):
    (path,) = export_csv([sample()], tmp_path / "new" / "folder")
    assert path.exists()

@pytest.mark.parametrize("text, written", [
    ("=1+1", "'=1+1"), ("@SUM(A1)", "'@SUM(A1)"), ("+cmd|' /C calc'!A0", "'+cmd|' /C calc'!A0"),
    ("-2+3", "'-2+3"), ("-abc", "'-abc"), ("\tTab", "'\tTab"),
    ("-3.4%", "-3.4%"), ("+5", "+5"), ("-1,200", "-1,200"), ("-", "-"), ("hello = world", "hello = world"),
])
def test_csv_stops_text_becoming_a_formula(tmp_path, text, written):
    (path,) = export_csv([Table("t", ["A"], [[text]])], tmp_path)
    with open(path, encoding="utf-8-sig", newline="") as handle:
        assert list(csv.reader(handle))[1] == [written]

def test_csv_formula_protection_can_be_switched_off(tmp_path):
    (path,) = export_csv([Table("t", ["A"], [["=1+1"]])], tmp_path, protect_formulas=False)
    with open(path, encoding="utf-8-sig", newline="") as handle:
        assert list(csv.reader(handle))[1] == ["=1+1"]

def test_csv_does_not_alter_real_numbers(tmp_path):
    (path,) = export_csv([Table("t", ["A"], [[-5], [-3.4], [0]])], tmp_path)
    with open(path, encoding="utf-8-sig", newline="") as handle:
        assert [r[0] for r in list(csv.reader(handle))[1:]] == ["-5", "-3.4", "0"]

#Excel
def test_xlsx_sheets_headers_and_numbers(tmp_path):
    path = export_xlsx([sample("Prices"), sample("Other: sheet")], tmp_path / "out.xlsx")
    book = load_workbook(path)
    assert book.sheetnames == ["Prices", "Other sheet"] # no empty default sheet
    sheet = book["Prices"]
    assert [c.value for c in sheet[1]] == ["Name", "Price"]
    assert sheet["B2"].value == 1234.5 and isinstance(sheet["B3"].value, int) # real numbers
    assert sheet["A2"].value == "Alpha"

def test_xlsx_formatting(tmp_path):
    sheet = load_workbook(export_xlsx([sample()], tmp_path / "out.xlsx"))["Prices"]
    assert sheet["A1"].font.bold and sheet["B1"].font.bold
    assert not sheet["A2"].font.bold
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B4"
    assert sheet.column_dimensions["A"].width >= 8

def test_xlsx_column_width_follows_the_content_but_is_capped(tmp_path):
    table = Table("t", ["Short", "Long"], [["a", "x" * 200], ["b", "y"]])
    sheet = load_workbook(export_xlsx([table], tmp_path / "w.xlsx"))["t"]
    assert sheet.column_dimensions["A"].width == 8
    assert sheet.column_dimensions["B"].width == 60

def test_xlsx_text_that_looks_like_a_formula_stays_text(tmp_path):
    table = Table("t", ["A"], [["=1+1"], ["@x"], ["+1"]])
    sheet = load_workbook(export_xlsx([table], tmp_path / "f.xlsx"))["t"]
    assert [sheet.cell(row=r, column=1).value for r in (2, 3, 4)] == ["=1+1", "@x", "+1"]
    assert all(sheet.cell(row=r, column=1).data_type == "s" for r in (2, 3, 4))

def test_xlsx_survives_control_characters_and_huge_text(tmp_path):
    table = Table("t", ["A\x07"], [["bad\x00\x01char"], ["z" * 40000]])
    sheet = load_workbook(export_xlsx([table], tmp_path / "c.xlsx"))["t"]
    assert sheet["A1"].value == "A"
    assert sheet["A2"].value == "badchar"
    assert len(sheet["A3"].value) == 32767

def test_xlsx_adds_the_extension_and_the_folder(tmp_path):
    path = export_xlsx([sample()], tmp_path / "deep" / "folder" / "report")
    assert path.name == "report.xlsx" and path.exists()

def test_xlsx_table_with_no_rows(tmp_path):
    sheet = load_workbook(export_xlsx([Table("t", ["A", "B"], [])], tmp_path / "e.xlsx"))["t"]
    assert [c.value for c in sheet[1]] == ["A", "B"]

def test_xlsx_refuses_a_table_too_big_for_excel(tmp_path, monkeypatch):
    monkeypatch.setattr(exporter, "XLSX_MAX_ROWS", 2)
    with pytest.raises(ExportError, match="CSV"):
        export_xlsx([Table("big", ["A"], [["1"], ["2"], ["3"]])], tmp_path / "x.xlsx")
    assert not (tmp_path / "x.xlsx").exists()

def test_xlsx_file_open_in_excel_gives_a_friendly_message(tmp_path, monkeypatch):
    def locked(self, filename):
        raise PermissionError(13, "Permission denied")
    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", locked)
    with pytest.raises(ExportError, match="open in another program"):
        export_xlsx([sample()], tmp_path / "x.xlsx")

#JSON
def test_json_structure(tmp_path):
    path = export_json([sample()], tmp_path / "out.json")
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["siphon_version"]
    (entry,) = data["tables"]
    assert entry["name"] == "Prices" and entry["source_url"] == "http://x.test/p"
    assert entry["columns"] == ["Name", "Price"]
    assert entry["rows"] == [{"Name": "Alpha", "Price": 1234.5}, {"Name": "Beta", "Price": 99}, {"Name": "Gamma", "Price": ""}]

def test_json_keeps_non_english_text_readable(tmp_path):
    path = export_json([Table("t", ["City"], [["\u0141\u00f3d\u017a"]])], tmp_path / "u.json")
    assert "\u0141\u00f3d\u017a" in path.read_text(encoding="utf-8")        # not \u0141\u00f3...

def test_json_repeated_column_names_do_not_lose_data(tmp_path):
    table = Table("t", ["A", "A"], [["1", "2"]])
    data = json.loads(export_json([table], tmp_path / "d.json").read_text(encoding="utf-8"))
    assert data["tables"][0]["rows"] == [{"A": "1", "A (2)": "2"}]

def test_json_adds_the_extension(tmp_path):
    assert export_json([sample()], tmp_path / "out").name == "out.json"

#shared behavior
@pytest.mark.parametrize("call", [
    lambda tmp: export_csv([], tmp),
    lambda tmp: export_xlsx([], tmp / "a.xlsx"),
    lambda tmp: export_json([], tmp / "a.json"),
])
def test_exporting_nothing_is_an_error_not_an_empty_file(tmp_path, call):
    with pytest.raises(ExportError, match="no tables"):
        call(tmp_path)
    assert list(tmp_path.iterdir()) == []

def test_a_folder_that_cannot_be_created_gives_a_friendly_message(tmp_path):
    blocker = tmp_path / "file.txt"
    blocker.write_text("i am a file, not a folder")
    with pytest.raises(ExportError, match="Couldn't create the folder"):
        export_csv([sample()], blocker / "sub")
