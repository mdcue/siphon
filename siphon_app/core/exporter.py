"""
siphon_app/core/exporter.py


Saves tables as CSV (one file each), Excel (one workbook, one sheet each) or JSON.
It never overwrites a file you didn't pick, strips characters Windows and Excel
refuse in names, and keeps scraped text from turning into spreadsheet formulas.
Nothing here knows about windows or buttons.
"""
import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from siphon_app import __version__
from siphon_app.core.models import Cell, Table

MAX_NAME = 100 # longest file name we create (without the extension)
SHEET_NAME_MAX = 31 # Excel's limit
XLSX_MAX_ROWS = 1_048_575 # Excel allows 1,048,576 rows; one is the header
XLSX_MAX_TEXT = 32_767 # Excel's limit for the text in one cell
MAX_COLUMN_WIDTH = 60

_RESERVED = {"con", "prn", "aux", "nul",
             *(f"com{i}" for i in range(1, 10)), *(f"lpt{i}" for i in range(1, 10))}

class ExportError(Exception):
    """Something went wrong while saving. The message is safe to show to the user."""

#names
def safe_filename(name: str) -> str:
    """'Population: 2020/21?' -> 'Population 2020 21'. Always returns something usable."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", name)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    cleaned = cleaned[:MAX_NAME].strip(" .")
    if not cleaned:
        return "table"
    if cleaned.split(".")[0].strip().lower() in _RESERVED: # windows won't allow "con.csv"
        cleaned = "_" + cleaned
    return cleaned

def _make_unique(names: list[str], limit: int) -> list[str]:
    """Add ' (2)', ' (3)' to repeats (ignoring capitals) without going over the limit."""
    seen: set[str] = set()
    result = []
    for name in names:
        candidate, count = name, 1
        while candidate.lower() in seen:
            count += 1
            suffix = f" ({count})"
            candidate = name[:limit - len(suffix)].rstrip() + suffix
        seen.add(candidate.lower())
        result.append(candidate)
    return result

def safe_sheet_names(names: list[str]) -> list[str]:
    """Excel sheet names: max 31 characters, none of [ ] : * ? / \\, unique, not 'History'."""
    cleaned = []
    for name in names:
        text = re.sub(r"[\[\]:*?/\\]", " ", name)
        text = re.sub(r"\s+", " ", text).strip().strip("'").strip()
        text = text[:SHEET_NAME_MAX].strip().strip("'").strip() or "Sheet"
        if text.lower() == "history": # reserved by Excel
            text += "_"
        cleaned.append(text)
    return _make_unique(cleaned, SHEET_NAME_MAX)

#helpers
def _os_problem(error: OSError, path: Path) -> ExportError:
    if isinstance(error, PermissionError):
        return ExportError(f"Couldn't write {path.name}. It may be open in another program "
                        "(like Excel): close it and try again.")
    return ExportError(f"Couldn't write {path.name}: {error.strerror or error}.")

def _free_path(folder: Path, name: str, extension: str) -> Path:
    """folder/name.csv, or 'name (2).csv' if that exists: never overwrite someone's file."""
    path = folder / f"{name}{extension}"
    count = 1
    while path.exists():
        count += 1
        path = folder / f"{name} ({count}){extension}"
    return path

def _check_tables(tables: list[Table]) -> None:
    if not tables:
        raise ExportError("There are no tables to export.")

def _protect(cell: Cell) -> Cell:
    """
    Stop a text cell being run as a formula when the CSV is opened in Excel.
    A cell starting with = or @, or with + / - that is not a plain number, gets a leading '.
    """
    if not isinstance(cell, str) or not cell:
        return cell
    first = cell[0]
    if first in "=@\t\r":
        return "'" + cell
    if first in "+-" and not re.fullmatch(r"[+-]\d[\d.,]*%?", cell) and len(cell) > 1:
        return "'" + cell
    return cell

#for CSVs
def export_csv(tables: list[Table], folder: str | Path, *, protect_formulas: bool = True) -> list[Path]:
    """
    One UTF-8 file per table (with a BOM so Excel shows accents and symbols properly).
    Returns the paths written. Existing files are never overwritten.
    """
    _check_tables(tables)
    folder = Path(folder)
    try:
        folder.mkdir(parents=True, exist_ok=True)
    except OSError as error:
        raise ExportError(f"Couldn't create the folder {folder}: {error.strerror or error}.") from None
    names = _make_unique([safe_filename(t.name) for t in tables], MAX_NAME)
    paths = []
    for table, name in zip(tables, names):
        path = _free_path(folder, name, ".csv")
        convert = _protect if protect_formulas else (lambda cell: cell)
        try:
            with open(path, "w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow([convert(h) for h in table.headers])
                writer.writerows([convert(cell) for cell in row] for row in table.rows)
        except OSError as error:
            raise _os_problem(error, path) from None
        paths.append(path)
    return paths

#for Excel
def _excel_text(text: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", text)[:XLSX_MAX_TEXT]

def export_xlsx(tables: list[Table], path: str | Path) -> Path:
    """
    One workbook, one sheet per table: bold frozen header, filter buttons, sensible
    column widths, numbers stored as numbers and text always stored as text.
    """
    _check_tables(tables)
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_name(path.name + ".xlsx")
    for table in tables:
        if table.n_rows > XLSX_MAX_ROWS:
            raise ExportError(f"'{table.name}' has {table.n_rows:,} rows, more than an Excel "
                            "sheet can hold. Export it as CSV instead.")

    workbook = Workbook()
    workbook.remove(workbook.active) # drop the empty default sheet
    for table, sheet_name in zip(tables, safe_sheet_names([t.name for t in tables])):
        sheet = workbook.create_sheet(sheet_name)
        widths = [len(h) for h in table.headers]
        for column, header in enumerate(table.headers, start=1):
            cell = sheet.cell(row=1, column=column, value=_excel_text(header))
            cell.data_type = "s"
            cell.font = Font(bold=True)
        for row_number, row in enumerate(table.rows, start=2):
            for column, value in enumerate(row, start=1):
                if isinstance(value, str):
                    value = _excel_text(value)
                cell = sheet.cell(row=row_number, column=column, value=value)
                if isinstance(value, str):
                    cell.data_type = "s" # "=1+1" stays text, not a formula
                if row_number <= 1000: # sample the first rows for widths
                    widths[column - 1] = max(widths[column - 1], len(str(value)))
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 8), MAX_COLUMN_WIDTH)
        if table.headers:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    except OSError as error:
        raise _os_problem(error, path) from None
    return path

#JSON
def export_json(tables: list[Table], path: str | Path) -> Path:
    """One JSON file: every table with its name, source and rows as {column: value} records."""
    _check_tables(tables)
    path = Path(path)
    if path.suffix.lower() != ".json":
        path = path.with_name(path.name + ".json")
    entries = []
    for table in tables:
        columns = _make_unique(list(table.headers), 10**6) # records can't repeat a key
        entries.append({
            "name": table.name,
            "source_url": table.source_url,
            "columns": columns,
            "rows": [dict(zip(columns, row)) for row in table.rows],
        })
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as handle:
            json.dump({"siphon_version": __version__, "tables": entries}, handle,
                    ensure_ascii=False, indent=2)
    except OSError as error:
        raise _os_problem(error, path) from None
    return path

if __name__ == "__main__":
    # Quick terminal check:  python -m siphon_app.core.exporter tests/fixtures/prices.html out
    import sys
    from siphon_app.core.cleaner import clean_table
    from siphon_app.core.parser import parse_tables

    if len(sys.argv) != 3:
        print("usage: python -m siphon_app.core.exporter <html file> <output folder>")
        raise SystemExit(1)
    with open(sys.argv[1], "rb") as handle:
        found = [clean_table(t)[0] for t in parse_tables(handle.read())]
    out = Path(sys.argv[2])
    for written in export_csv(found, out) + [export_xlsx(found, out / "tables"), export_json(found, out / "tables")]:
        print("wrote", written)